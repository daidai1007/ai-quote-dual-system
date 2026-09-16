import json
import os
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import Image


source = Path(os.environ["ATTACHMENT_DOCX_PATH"])
output = Path(os.environ["ATTACHMENT_DOCX_REPORT"])
document = Document(source)


def images_in(element):
    found = []
    for blip in element.xpath(".//a:blip"):
        relation_id = blip.get(qn("r:embed"))
        if not relation_id:
            continue
        part = document.part.related_parts.get(relation_id)
        if part is None:
            continue
        dimensions = None
        try:
            from io import BytesIO

            with Image.open(BytesIO(part.blob)) as image:
                dimensions = [image.width, image.height]
        except Exception:
            pass
        found.append(
            {
                "relation_id": relation_id,
                "partname": str(part.partname),
                "bytes": len(part.blob),
                "dimensions": dimensions,
            }
        )
    return found


payload = {
    "source": str(source),
    "paragraphs": [
        {"index": index, "text": paragraph.text, "images": images_in(paragraph._p)}
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text.strip() or images_in(paragraph._p)
    ],
    "tables": [],
}
for table_index, table in enumerate(document.tables):
    rows = []
    for row_index, row in enumerate(table.rows):
        cells = []
        for column_index, cell in enumerate(row.cells):
            cells.append(
                {
                    "column": column_index,
                    "text": "\n".join(
                        paragraph.text for paragraph in cell.paragraphs if paragraph.text.strip()
                    ),
                    "images": images_in(cell._tc),
                }
            )
        rows.append({"row": row_index, "cells": cells})
    payload["tables"].append({"table": table_index, "rows": rows})

output.parent.mkdir(parents=True, exist_ok=True)
output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({"paragraphs": len(payload["paragraphs"]), "tables": len(payload["tables"]), "output": str(output)}, ensure_ascii=False))
